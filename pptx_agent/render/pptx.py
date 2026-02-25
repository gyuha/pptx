from __future__ import annotations

# pyright: reportAttributeAccessIssue=false, reportMissingImports=false, reportMissingTypeStubs=false, reportUnknownArgumentType=false, reportUnknownMemberType=false, reportUnknownParameterType=false, reportUnknownVariableType=false

import hashlib
from pathlib import Path
from typing import TypedDict

from pptx.api import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..contracts import Slide, SlideElement, SlideSpec
from .image_api import (
    ImageGenerationFailedError,
    generate_image_bytes,
    resolve_default_image_model,
)

SUPPORTED_LAYOUT_HINTS = {"title_body_flow", "title_table_focus"}

SLIDE_WIDTH_IN = 10.0
SLIDE_HEIGHT_IN = 7.5

TITLE_X_IN = 0.5
TITLE_Y_IN = 0.3
TITLE_W_IN = 9.0
TITLE_H_IN = 1.0

CONTENT_X_IN = 0.7
CONTENT_Y_IN = 1.5
CONTENT_W_IN = 8.6
CONTENT_H_IN = 5.5
CONTENT_GAP_IN = 0.2

BODY_MIN_H_IN = 0.7
TABLE_MIN_H_IN = 1.2

TITLE_MAX_CHARS = 140
BODY_MAX_CHARS = 420
TABLE_CELL_MAX_CHARS = 120


class UnsupportedLayoutHintError(ValueError):
    pass


class ImageTrace(TypedDict):
    slide_id: str
    element_index: int
    image_prompt: str
    image_model: str
    image_seed: int | None
    image_path: str


class RenderImageGenerationStats(TypedDict):
    requested: int
    succeeded: int
    failed: int
    image_model: str | None


class RenderQAReport(TypedDict):
    deck_id: str
    schema_version: str
    output_pptx: str
    qa_schema: str
    slide_count_expected: int
    slide_count_rendered: int
    critical_overflow_count: int
    overflow_metric_note: str
    image_generation: RenderImageGenerationStats
    image_traces: list[ImageTrace]


def _layout_hint_for_slide(slide: Slide) -> str:
    if slide.layout_hint is None:
        if any(element.kind == "table" for element in slide.elements):
            return "title_table_focus"
        return "title_body_flow"
    if slide.layout_hint not in SUPPORTED_LAYOUT_HINTS:
        raise UnsupportedLayoutHintError(
            f"unsupported-layout hint='{slide.layout_hint}' on slide_id='{slide.slide_id}'"
        )
    return slide.layout_hint


def _set_text_frame(
    text_frame: object,
    text: str,
    *,
    font_size_pt: int,
    bold: bool,
    max_chars: int,
) -> int:
    frame = text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(4)
    frame.margin_right = Pt(4)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.level = 0
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_after = Pt(6)

    run = paragraph.runs[0]
    run.font.bold = bold
    run.font.size = Pt(font_size_pt)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    return 1 if len(text) > max_chars else 0


def _split_table_text(raw_text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in raw_text.splitlines():
        clean_line = line.strip()
        if not clean_line:
            continue
        if "\t" in clean_line:
            cells = [cell.strip() for cell in clean_line.split("\t")]
        else:
            cells = [clean_line]
        rows.append(cells)
    return rows


def _render_table_or_fallback(
    pptx_slide: object,
    element: SlideElement,
    *,
    top_in: float,
    height_in: float,
) -> tuple[int, float]:
    rows = _split_table_text(element.text or "")
    if not rows:
        text_box = pptx_slide.shapes.add_textbox(
            Inches(CONTENT_X_IN),
            Inches(top_in),
            Inches(CONTENT_W_IN),
            Inches(height_in),
        )
        overflow = _set_text_frame(
            text_box.text_frame,
            element.text or "",
            font_size_pt=18,
            bold=False,
            max_chars=BODY_MAX_CHARS,
        )
        return overflow, top_in + height_in + CONTENT_GAP_IN

    max_cols = max(len(row) for row in rows)
    if max_cols <= 0 or len(rows) > 12 or max_cols > 8:
        text_box = pptx_slide.shapes.add_textbox(
            Inches(CONTENT_X_IN),
            Inches(top_in),
            Inches(CONTENT_W_IN),
            Inches(height_in),
        )
        overflow = _set_text_frame(
            text_box.text_frame,
            element.text or "",
            font_size_pt=18,
            bold=False,
            max_chars=BODY_MAX_CHARS,
        )
        return overflow, top_in + height_in + CONTENT_GAP_IN

    table_shape = pptx_slide.shapes.add_table(
        rows=len(rows),
        cols=max_cols,
        left=Inches(CONTENT_X_IN),
        top=Inches(top_in),
        width=Inches(CONTENT_W_IN),
        height=Inches(height_in),
    )
    table = table_shape.table

    overflow_count = 0
    for row_index, row in enumerate(rows):
        for col_index in range(max_cols):
            cell_text = row[col_index] if col_index < len(row) else ""
            cell = table.cell(row_index, col_index)
            cell.text = cell_text
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.bold = row_index == 0
            overflow_count += 1 if len(cell_text) > TABLE_CELL_MAX_CHARS else 0

    return overflow_count, top_in + height_in + CONTENT_GAP_IN


def _render_image(
    *,
    pptx_slide: object,
    slide: Slide,
    element: SlideElement,
    element_index: int,
    top_in: float,
    height_in: float,
    image_assets_dir: Path,
    default_image_model: str,
) -> tuple[ImageTrace, float]:
    prompt = (
        element.image_prompt
        if element.image_prompt is not None
        else (element.text if element.text is not None else slide.title)
    )
    prompt = prompt.strip()
    if not prompt:
        raise ImageGenerationFailedError(
            f"image prompt must not be blank for slide_id='{slide.slide_id}'"
        )
    model = (
        element.image_model.strip()
        if element.image_model is not None
        else default_image_model
    )
    if not model:
        raise ImageGenerationFailedError(
            f"image model must not be blank for slide_id='{slide.slide_id}'"
        )
    seed = element.image_seed

    image_bytes = generate_image_bytes(prompt=prompt, model=model, seed=seed)
    digest_source = f"{slide.slide_id}|{element_index}|{prompt}|{model}|{seed}"
    digest = hashlib.sha256(digest_source.encode("utf-8")).hexdigest()[:12]
    file_name = f"{slide.slide_id}_image_{element_index:02d}_{digest}.png"
    image_path = image_assets_dir / file_name
    image_path.parent.mkdir(parents=True, exist_ok=True)
    _ = image_path.write_bytes(image_bytes)

    _ = pptx_slide.shapes.add_picture(
        str(image_path),
        Inches(CONTENT_X_IN),
        Inches(top_in),
        width=Inches(CONTENT_W_IN),
        height=Inches(height_in),
    )

    trace: ImageTrace = {
        "slide_id": slide.slide_id,
        "element_index": element_index,
        "image_prompt": prompt,
        "image_model": model,
        "image_seed": seed,
        "image_path": str(image_path),
    }
    return trace, top_in + height_in + CONTENT_GAP_IN


def render_slidespec_to_pptx(slidespec: SlideSpec, output_path: Path) -> RenderQAReport:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)

    critical_overflow_count = 0
    image_traces: list[ImageTrace] = []
    image_requested = 0
    image_succeeded = 0
    image_failed = 0
    default_image_model = resolve_default_image_model()
    image_assets_dir = output_path.with_suffix(output_path.suffix + ".images")

    blank_layout = presentation.slide_layouts[6]
    for slide in slidespec.slides:
        layout_hint = _layout_hint_for_slide(slide)
        pptx_slide = presentation.slides.add_slide(blank_layout)

        title_box = pptx_slide.shapes.add_textbox(
            Inches(TITLE_X_IN),
            Inches(TITLE_Y_IN),
            Inches(TITLE_W_IN),
            Inches(TITLE_H_IN),
        )
        critical_overflow_count += _set_text_frame(
            title_box.text_frame,
            slide.title,
            font_size_pt=34,
            bold=True,
            max_chars=TITLE_MAX_CHARS,
        )

        body_elements = [
            element
            for element in slide.elements
            if element.kind in {"body", "table", "title", "image"}
        ]
        if body_elements and body_elements[0].kind == "title":
            body_elements = body_elements[1:]

        if not body_elements:
            continue

        if layout_hint == "title_table_focus":
            table_weight = 2.0
            body_weight = 1.0
            image_weight = 1.2
        else:
            table_weight = 1.3
            body_weight = 1.0
            image_weight = 1.6

        weights = [
            table_weight
            if element.kind == "table"
            else (image_weight if element.kind == "image" else body_weight)
            for element in body_elements
        ]
        total_weight = sum(weights)
        available_height = CONTENT_H_IN - (
            max(0, len(body_elements) - 1) * CONTENT_GAP_IN
        )

        current_top = CONTENT_Y_IN
        for index, element in enumerate(body_elements):
            height_share = available_height * (weights[index] / total_weight)
            min_height = TABLE_MIN_H_IN if element.kind == "table" else BODY_MIN_H_IN
            block_height = max(min_height, round(height_share, 3))
            if index == len(body_elements) - 1:
                block_height = max(
                    min_height,
                    round((CONTENT_Y_IN + CONTENT_H_IN) - current_top, 3),
                )

            if element.kind == "table":
                overflow_delta, next_top = _render_table_or_fallback(
                    pptx_slide,
                    element,
                    top_in=current_top,
                    height_in=block_height,
                )
                critical_overflow_count += overflow_delta
                current_top = next_top
                continue

            if element.kind == "image":
                image_requested += 1
                try:
                    trace, next_top = _render_image(
                        pptx_slide=pptx_slide,
                        slide=slide,
                        element=element,
                        element_index=index,
                        top_in=current_top,
                        height_in=block_height,
                        image_assets_dir=image_assets_dir,
                        default_image_model=default_image_model,
                    )
                except ImageGenerationFailedError:
                    image_failed += 1
                    raise
                image_traces.append(trace)
                image_succeeded += 1
                current_top = next_top
                continue

            text_box = pptx_slide.shapes.add_textbox(
                Inches(CONTENT_X_IN),
                Inches(current_top),
                Inches(CONTENT_W_IN),
                Inches(block_height),
            )
            critical_overflow_count += _set_text_frame(
                text_box.text_frame,
                element.text or "",
                font_size_pt=21,
                bold=False,
                max_chars=BODY_MAX_CHARS,
            )
            current_top += block_height + CONTENT_GAP_IN

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))

    return {
        "deck_id": slidespec.deck_id,
        "schema_version": slidespec.schema_version,
        "output_pptx": str(output_path),
        "qa_schema": "render.qa.v1",
        "slide_count_expected": len(slidespec.slides),
        "slide_count_rendered": len(presentation.slides),
        "critical_overflow_count": critical_overflow_count,
        "overflow_metric_note": (
            "deterministic_heuristic_char_thresholds:title=140,body=420,table_cell=120"
        ),
        "image_generation": {
            "requested": image_requested,
            "succeeded": image_succeeded,
            "failed": image_failed,
            "image_model": image_traces[0]["image_model"] if image_traces else None,
        },
        "image_traces": image_traces,
    }
